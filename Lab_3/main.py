import matplotlib.pyplot as plt
import numpy as np
import mplcursors
import pandas as pd
import docx as docx

def statCalulator(array, type):
    match type:
        # середнє значення
        case "average":
            return round(np.mean(array), 8)
        # середнє квадратичне відхилення
        case "std":
            return round(np.std(array), 8)
        # максимальне відхилення від середнього значення
        case "maxAverageDeviation":
            return round(max(list(map(lambda x: x != 0 and abs(np.mean(array) - x), array))), 8)


def ticksSeparator(timeArray, dataArray, ticks):
    result = []
    result.append([timeArray[0], dataArray[0]])
    for i in range(ticks,len(timeArray)-1,ticks):
        tmpRes = []
        for j in range(i-ticks, i):
            if timeArray[i] > (timeArray[i - ticks] + ticks):
                dataArray[i] = (dataArray[i-1]+dataArray[i+1])/2
            elif timeArray[i] < (timeArray[i - ticks] + ticks):
                dataArray[i] = (dataArray[i] + dataArray[i + 1]) / 2
            tmpRes.append(dataArray[j])
        tmp= np.mean(tmpRes)
        result.append([timeArray[i], tmp])
    return np.array(result)

def chartBuilder(xData, yData, title):
    average = statCalulator(yData, 'average')
    maxAverDiv = statCalulator(yData, 'maxAverageDeviation')
    def textRun(annotation):
        syt = annotation.replace('x', "seconds")
        syt = syt.replace('y', "degree")
        dd = list(map(str, syt.split('=')))
        valueCon = None
        try:
            valueCon = float(dd[len(dd) - 1])
        except ValueError:
            valueCon = float(dd[len(dd) - 1][1:]) * (-1)
        diff = abs(average - valueCon)
        syt = syt + "\n deviation " + str(diff)
        return syt
    averageArray = np.full(len(xData),average)
    plt.style.use('seaborn-v0_8')
    fig, ax = plt.subplots()
    fig.set(facecolor='lightblue')
    ax.plot(xData, yData, color="blue")
    ax.plot(xData, averageArray, color='red')
    ax.set_xlabel("seconds", fontsize=20)
    ax.set_ylabel("degree", fontsize=20)
    plt.title(title + " average value: " + str(round(average, 8)), fontsize=20)
    dots = ax.scatter(xData, yData, color='darkred')
    crs = mplcursors.cursor(dots, hover=True)
    crs.connect("add", lambda sel: sel.annotation.set_text(textRun(sel.annotation.get_text())))
    for x, y in zip(xData, yData):

        curPoint = (float(abs(float(y) - average)))
        if curPoint == maxAverDiv:
            ax.scatter(x, y, color='red')

    plt.legend(['dependency', 'Average', 'checkpoint', 'maxDif'], loc='center left', bbox_to_anchor=(1.0, 0.5),
               fontsize=12, frameon=True)
    plt.show()

def genearateTable(headArr, pitchArr, timeArr, fileName):
    doc = docx.Document()

    parameters = ['Середнє значення (head)',
                  'Середнє значення (pitch)',
                  'Максимальне відхилення від середнього значення (head)',
                  'Максимальне відхилення від середнього значення (pitch)',
                  'Середньо квадратичне відхилення (head)',
                  'Середньо квадратичне відхилення (pitch)',
                  'Середнє значення з кроком 10с (head)',
                  'Середнє значення з кроком 10с (pitch)',
                  'Максимальне відхилення від середнього значення з кроком 10с (head)',
                  'Максимальне відхилення від середнього значення з кроком 10с (pitch)',
                  'Середньо квадратичне відхилення з кроком 10с (head)',
                  'Середньо квадратичне відхилення з кроком 10с (pitch)',
                  'Середнє значення з кроком 20с (head)',
                  'Середнє значення з кроком 20с (pitch)',
                  'Максимальне відхилення від середнього значення з кроком 20с (head)',
                  'Максимальне відхилення від середнього значення з кроком 20с (pitch)',
                  'Середньо квадратичне відхилення з кроком 20с (head)',
                  'Середньо квадратичне відхилення з кроком 20с (pitch)'
                  ]

    values = [statCalulator(headArr, 'average'),
              statCalulator(pitchArr, 'average'),
              statCalulator(headArr, 'maxAverageDeviation'),
              statCalulator(pitchArr, 'maxAverageDeviation'),
              statCalulator(headArr, 'std'),
              statCalulator(pitchArr, 'std'),
              statCalulator((ticksSeparator(timeArr, headArr, 10)[:, 1]), 'average'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 10)[:, 1]), 'average'),
              statCalulator((ticksSeparator(timeArr, headArr, 10)[:, 1]), 'maxAverageDeviation'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 10)[:, 1]), 'maxAverageDeviation'),
              statCalulator((ticksSeparator(timeArr, headArr, 10)[:, 1]), 'std'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 10)[:, 1]), 'std'),
              statCalulator((ticksSeparator(timeArr, headArr, 20)[:, 1]), 'average'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 20)[:, 1]), 'average'),
              statCalulator((ticksSeparator(timeArr, headArr, 20)[:, 1]), 'maxAverageDeviation'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 20)[:, 1]), 'maxAverageDeviation'),
              statCalulator((ticksSeparator(timeArr, headArr, 20)[:, 1]), 'std'),
              statCalulator((ticksSeparator(timeArr, pitchArr, 20)[:, 1]), 'std')
              ]

    index = np.arange(1, np.array(values).size + 1)

    table_data = {'Index': index, 'Parameter': parameters, 'Value': values}

    df = pd.DataFrame(data=table_data)

    # Initialise the table
    t = doc.add_table(rows=(df.shape[0] + 1), cols=df.shape[1])
    t.style = 'Light Shading Accent 1'

    # Add the column headings
    for j in range(df.shape[1]):
        t.cell(0, j).text = df.columns[j]

    # Add the body of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            cell = df.iat[i, j]
            t.cell(i + 1, j).text = str(cell)

    doc.save(fileName + '.docx')


dataGPS = []
d = {}
gpsData = open('gpsData.csv', 'r')

for line in gpsData.readlines():
    s = list(map(str, line.split(',')))
    dataGPS.append(s)

def valueParser(columnNumber):

    tmp = []
    for i in range(len(dataGPS)):
        tmp.append(dataGPS[i][columnNumber - 1])
    return np.array(tmp, dtype=np.float64)

# Get values according to the variant
d['head'] = valueParser(13)
d['pitch'] = valueParser(14)
d['time'] = valueParser(7)


chartBuilder(d['time'], d['head'], 'Dependence of time on (head) degrees')
chartBuilder(d['time'], d['pitch'], 'Dependence of time on (pitch) degrees')

chartBuilder((ticksSeparator(d['time'], d['head'], 10)[:, 0]), (ticksSeparator(d['time'], d['head'], 10)[:, 1]), 'Dependence of time on (head) degrees 10s intervel')
chartBuilder((ticksSeparator(d['time'], d['pitch'], 10)[:, 0]), (ticksSeparator(d['time'], d['pitch'], 10)[:, 1]), 'Dependence of time on (pitch) degrees 10s intervel')

chartBuilder((ticksSeparator(d['time'], d['head'], 20)[:, 0]), (ticksSeparator(d['time'], d['head'], 20)[:, 1]), 'Dependence of time on (head) degrees 20s intervel')
chartBuilder((ticksSeparator(d['time'], d['pitch'], 20)[:, 0]), (ticksSeparator(d['time'], d['pitch'], 20)[:, 1]), 'Dependence of time on (pitch) degrees 20s intervel')

genearateTable(d['head'], d['pitch'], d['time'], 'Results')
