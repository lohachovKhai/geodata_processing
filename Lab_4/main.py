import numpy as np
import docx as docx
import statistics as stat
import pandas as pd
import random
from fractions import Fraction as F
from decimal import Decimal as D

def genearateTable(course, velocity, time, fileName):
    doc = docx.Document()

    types= ['Істинний курс, градус', 'Швидкість польоту, км/год', 'Час руху ЛА, хв']
    meanArr=[(stat.mean(course)), (stat.mean(velocity)), (stat.mean(time))]
    modeArr=[(stat.mode(course)), (stat.mode(velocity)), (stat.mode(time))]
    medianArr=[(stat.median(course)), (stat.median(velocity)), (stat.median(time))]
    stdevArr=[(stat.stdev(course)), (stat.stdev(velocity)), (stat.stdev(time))]
    pvarianceArr=[(stat.pvariance(course)), (stat.pvariance(velocity)), (stat.pvariance(time))]

    index = np.arange(1, 4)

    table_data = {'№ п/п': index, '': types, 'Середнє значення': meanArr, 'Мода': modeArr,
                  'Медіана': medianArr, 'Сер.квадр. відхилення': np.round(stdevArr, decimals=3),
                  'Дисперсія': pvarianceArr}

    df = pd.DataFrame(data=table_data)

    # Initialise the table
    t = doc.add_table(rows=(df.shape[0] + 1), cols=df.shape[1])
    t.style = 'Light Shading Accent 1'

    # Add the column headings
    for j in range(df.shape[1]):
        t.cell(0, j).text = df.columns[j]

    # Add the body of the data frame
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            cell = df.iat[i, j]
            t.cell(i + 1, j).text = str(cell)

    doc.save(fileName + '.docx')

data = {
    'Stage': np.arange(1, 11),
    'Course':   [100, 280, 48, 50, 100, 5, 16, 25, 44, 16],
    'Velocity': [36, 250, 58, 100, 16, 163, 240, 58, 165, 50],
    'Time':     [12, 36, 24, 5, 8, 16, 18, 22, 6, 16]}

genearateTable(data['Course'], data['Velocity'], data['Time'], 'Result')


examplesArr = [
    (1, stat.mean([5, 51, 44, 99, 1])),
    (2, stat.mean([F(36, 24), F(12, 13), F(50, 90), F(7, 19)])),
    (3, stat.mean([D("12.5"), D("1.15"), D("15.1"), D("22.25")])),
    (4, stat.mean([ random.randint(1, 500) for x in range(1,100000) ])),
    (5, stat.mean([random.triangular(1, 500, 80) for x in range(1,100000) ])),
    (6, stat.mode([ random.randint(1, 500) for x in range(1,100000)])),
    (7, stat.mode([ random.randint(1, 500) for x in range(1,100000) ])),
    (8, stat.mode([ random.randint(1, 500) for x in range(1,100000) ])),
    (9, stat.mode(["cat", "dog", "dog", "monkey", "cat", "monkey", "dog"])),
    (10, stat.median([ random.randint(1, 500) for x in range(1,100000) ])),
    (11, stat.median([ random.randint(1, 500) for x in range(1,100000) ])),
    (12, stat.median([ random.randint(1, 500) for x in range(1,100000) ])),
    (13, stat.median_high([ random.randint(1, 500) for x in range(1,100000) ])),
    (14, stat.median([ random.randint(1, 500) for x in range(1,100000) ])),
    (15, stat.pvariance([3, 4, 5, 5, 5, 5, 5, 6, 6])),
    (16, stat.pstdev([3, 4, 5, 5, 5, 5, 5, 6, 6])),
    (17, stat.variance([3, 4, 5, 5, 5, 5, 5, 6, 6])),
    (18, stat.pvariance([21, 22, 23, 24, 25, 26, 27, 28, 29])),
    (19, stat.pstdev([7, 8, 9, 9, 9, 9, 9, 4, 4])),
    (20, stat.variance([F(77, 88), F(74, 152), F(155, 212)]))
]

print('\n'.join('     '.join(str(item)+(" " if len(str(item))==1 else "") for item in tup) for tup in examplesArr))